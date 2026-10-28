"""Generate the synthetic sample documents in samples/ (all content is invented).

    python samples/make_samples.py

* AMM_32-42-27_brake_wear_check_rev61.pdf  born-digital PDF with a text layer   -> fast lane
* TaskCard_VT-EXA_C-check_scan.pdf          image-only PDF (no text layer)        -> heavy lane (OCR)
* SOP_cold_weather_operations.docx           Word SOP with a table                  -> fast lane
* SB-A320-32-1471_summary.txt                plain text service bulletin summary    -> fast lane
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent

AMM_TASK = [
    ("h", "AIRCRAFT MAINTENANCE MANUAL - A320NEO (SYNTHETIC)"),
    ("h", "TASK 32-42-27-200-801-A Brake Wear Check"),
    (
        "p",
        "Effectivity: ALL. Revision 61. This task gives the procedure to check the wear of the "
        "main landing gear carbon brakes with the wear indicator pins. Reference ATA 32-42-27.",
    ),
    ("h", "1. Reason for the Job"),
    (
        "p",
        "Do a check of the brake wear at each daily check and when the brake temperature is "
        "less than 150 deg.C. Use the wear pin method only with the parking brake applied.",
    ),
    ("h", "2. Job Set-up Information"),
    (
        "p",
        "A. Fixtures, Tools, Test and Support Equipment: No specific tooling. P/N 98D32103500-000 "
        "wear gauge may be used for the dimension check.",
    ),
    (
        "w",
        "WARNING: MAKE SURE THAT THE PARKING BRAKE IS SET AND THE WHEEL CHOCKS ARE IN POSITION. "
        "THE BRAKES CAN BE HOT AND CAN CAUSE BURNS.",
    ),
    (
        "p",
        "(1) Apply the parking brake and make sure that the brake pressure indication shows at least "
        "2000 psi on the triple indicator.",
    ),
    ("h", "3. Procedure"),
    (
        "p",
        "(1) Measure the length of the wear indicator pin that extends from the housing face. "
        "(2) If the pin end is flush with the housing face or below it, replace the brake assembly "
        "per AMM TASK 32-42-27-000-801. (3) If the pin extends more than 0.0 mm but less than 1.0 mm, "
        "record the measurement in the technical log and check again at the next daily check. "
        "(4) Make an entry in the aircraft technical log for VT-EXA or the applicable registration.",
    ),
    (
        "c",
        "CAUTION: DO NOT USE THE WEAR PIN READING WITH THE PARKING BRAKE RELEASED. THE READING WILL BE "
        "INCORRECT.",
    ),
    ("p", "(5) Close-up: remove the wheel chocks only after the brake temperature check is complete."),
]


def make_amm_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    for page in range(3):
        for kind, text in AMM_TASK:
            style = styles["Heading3"] if kind == "h" else styles["BodyText"]
            story.append(Paragraph(text if page == 0 else text.replace("801-A", f"801-{'ABC'[page]}"), style))
            story.append(Spacer(1, 6))
        story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=A4, title="AMM 32-42-27 (synthetic)").build(story)


def make_scanned_pdf(path: Path) -> None:
    """Draws only lines and boxes (no text objects): looks like a scanned form to the router."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for _ in range(4):
        for row in range(25):
            c.rect(50, 780 - row * 28, 495, 22)
            c.line(300, 780 - row * 28, 300, 802 - row * 28)
        c.showPage()
    c.save()


def make_sop_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("STANDARD OPERATING PROCEDURE - COLD WEATHER OPERATIONS (SYNTHETIC)", 1)
    d.add_paragraph("Applies to A320neo and ATR 72 fleets at stations with forecast OAT below +5 deg.C.")
    d.add_paragraph("NOTE: De-icing fluid holdover times start at the beginning of the final application.")
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate(
        [
            ("Fluid", "Mix", "Holdover (light snow)"),
            ("Type I", "50/50", "6-11 min"),
            ("Type IV", "100/0", "35-75 min"),
        ]
    ):
        for col, value in enumerate(row):
            table.cell(r, col).text = value
    d.add_paragraph("1. The captain confirms the holdover time with the de-icing crew before taxi.")
    d.save(str(path))


def make_sb_txt(path: Path) -> None:
    path.write_text(
        "SERVICE BULLETIN SB A320-32-1471 Rev 02 (SYNTHETIC SUMMARY)\n"
        "Subject: Main landing gear - brake wear indicator pin - replacement with improved pin.\n"
        "Effectivity: A320neo family MSN 10001-10999.\n"
        "Reason: Reports of wear pins bending during wheel change, giving a wrong wear reading.\n"
        "Compliance: Recommended at next wheel change. Reference AMM TASK 32-42-27-200-801.\n"
        "Parts: P/N C20195000-2 replaces P/N C20195000-1.\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    make_amm_pdf(HERE / "AMM_32-42-27_brake_wear_check_rev61.pdf")
    make_scanned_pdf(HERE / "TaskCard_VT-EXA_C-check_scan.pdf")
    make_sop_docx(HERE / "SOP_cold_weather_operations.docx")
    make_sb_txt(HERE / "SB-A320-32-1471_summary.txt")
    print("samples written to", HERE)
